from pathlib import Path

import pytest
from docx import Document

from proofreader.parsers.docx_parser import (
    DocumentSectionType,
    DocumentType,
    ParagraphType,
    TableType,
    classify_paragraph,
    classify_table,
    infer_heading_level,
    infer_section_type,
    parse_docx_with_sections,
)


def test_infer_heading_level_chinese():
    assert infer_heading_level("一、项目概述") == 1
    assert infer_heading_level("1.1 基础运维服务") == 2
    assert infer_heading_level("1.1.1 桌面运维") == 3
    assert infer_heading_level("(1) 工作内容") == 3


def test_classify_paragraph():
    assert classify_paragraph("1. 服务期限：12个月") == ParagraphType.NUMBERED_REQUIREMENT
    assert classify_paragraph("详见招标文件") == ParagraphType.PLAIN_TEXT


def test_infer_section_type_requirements():
    assert infer_section_type("第二部分 采购需求") == DocumentSectionType.REQUIREMENTS
    assert infer_section_type("用户需求书") == DocumentSectionType.REQUIREMENTS


def test_infer_section_type_bid_template():
    assert infer_section_type("第五部分 投标文件格式") == DocumentSectionType.BID_TEMPLATE
    assert infer_section_type("响应文件格式") == DocumentSectionType.BID_TEMPLATE


def test_classify_table_technical_spec():
    assert classify_table(["指标项", "技术要求"]) == TableType.TECHNICAL_SPEC


def test_classify_table_service_list():
    assert classify_table(["序号", "服务项", "服务频率", "服务要求"]) == TableType.SERVICE_LIST


def test_classify_table_evaluation():
    assert classify_table(["评审因素", "评分标准"]) == TableType.EVALUATION


def test_classify_table_unknown():
    assert classify_table(["foo", "bar"]) == TableType.UNKNOWN


@pytest.fixture
def sample_docx():
    return Path(__file__).resolve().parents[1] / "data" / "sample-docs" / "requirements.docx"


def test_parse_docx_with_sections_returns_document(sample_docx):
    result = parse_docx_with_sections(sample_docx)
    assert result.path == sample_docx
    assert result.doc_type == DocumentType.TENDER
    assert result.sections
    assert result.blocks
    assert all(block.paragraph_type is not None for block in result.blocks)


def test_parse_docx_with_sections_last_section_appended(sample_docx):
    result = parse_docx_with_sections(sample_docx)
    assert result.sections[-1].end_index > 0 or len(result.sections) == 1


def _make_sectioned_docx_with_table(path: Path) -> Path:
    """生成一个多 section 文档：表格出现在 REQUIREMENTS section 下。"""
    doc = Document()
    doc.add_heading("一、项目概述", level=1)
    doc.add_paragraph("项目概述内容。")
    doc.add_heading("二、采购需求", level=1)
    doc.add_paragraph("需求说明。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标项"
    table.cell(0, 1).text = "技术要求"
    table.cell(1, 0).text = "响应时间"
    table.cell(1, 1).text = "1小时"
    doc.add_heading("三、投标文件格式", level=1)
    doc.add_paragraph("请按格式提交。")
    doc.save(str(path))
    return path


def test_parse_docx_with_sections_table_attached_to_active_section(tmp_path):
    docx_path = _make_sectioned_docx_with_table(tmp_path / "sectioned.docx")
    result = parse_docx_with_sections(docx_path)

    # 找到采购需求与投标文件格式两个 section
    req_section = next(
        s for s in result.sections if s.section_type == DocumentSectionType.REQUIREMENTS
    )
    bid_section = next(
        s for s in result.sections if s.section_type == DocumentSectionType.BID_TEMPLATE
    )

    assert len(req_section.tables) == 1
    assert req_section.tables[0].table_type == TableType.TECHNICAL_SPEC.value
    assert len(bid_section.tables) == 0
    assert len(result.sections[-1].tables) == 0
    assert len(result.raw_tables) == 1
